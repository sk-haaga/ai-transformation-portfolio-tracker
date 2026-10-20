"""Build the one-slide leadership brief with python-pptx.

Layout (16:9):
    header        title, period, assessed RAG counts
    headline      the single sentence leadership must hear
    left panel    every initiative: assessed RAG, % complete, most at-risk milestone
    right panel   top 5 portfolio risks ranked by impact, with who they hit
    bottom strip  decisions required
    notes         key messages, full risk register, data-quality warnings
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .analysis import PortfolioAnalysis
from .registry import Registry
from .schema import InitiativeExtraction, Milestone, PortfolioSynthesis, RankedPortfolioRisk

NAVY = RGBColor(0x14, 0x23, 0x3B)
INK = RGBColor(0x26, 0x2B, 0x33)
GREY = RGBColor(0x5F, 0x66, 0x70)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD9, 0xDD, 0xE3)
ACCENT = RGBColor(0x2F, 0x5D, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RAG_COLOURS = {
    "RED": RGBColor(0xC6, 0x28, 0x28),
    "AMBER": RGBColor(0xE0, 0x9A, 0x00),
    "GREEN": RGBColor(0x2E, 0x7D, 0x32),
}
FONT = "Calibri"
TOP_RISKS = 5


# ------------------------------------------------------------------ helpers

def _text(slide, x, y, w, h, text, size=11, bold=False, colour=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    runs = text if isinstance(text, list) else [(text, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for chunk, style in runs:
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.name = FONT
        f.size = Pt(style.get("size", size))
        f.bold = style.get("bold", bold)
        f.italic = style.get("italic", italic)
        f.color.rgb = style.get("colour", colour)
    return box


def _shape(slide, kind, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:  # drop theme style refs so no renderer adds a shadow or bevel
        s._element.remove(style)
    return s


def _label_in(shape, text, size, colour=WHITE, bold=True):
    tf = shape.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = colour


def _rule(slide, x, y, w):
    line = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    line.line.color.rgb = RULE
    line.line.width = Pt(0.75)


def _fmt(d: date | None) -> str:
    return f"{d.day} {d:%b}" if d else "TBC"


def _slip(m: Milestone) -> int | None:
    return (m.forecast - m.baseline).days if m.forecast and m.baseline else None


def watch_milestone(e: InitiativeExtraction) -> Milestone | None:
    """The milestone leadership should watch.

    Largest known slip among late / at-risk milestones; if none has a known slip,
    the earliest flagged one; if nothing is flagged, the next open milestone.
    """
    open_ms = [m for m in e.milestones if m.status != "done"]
    flagged = [m for m in open_ms if m.status in ("late", "at_risk")]
    slipped = [m for m in flagged if (_slip(m) or 0) > 0]
    if slipped:
        return max(slipped, key=lambda m: (_slip(m), -(m.baseline or date.max).toordinal()))
    pool = flagged or open_ms
    return min(pool, key=lambda m: m.baseline or m.forecast or date.max) if pool else None


def milestone_text(m: Milestone | None) -> list[tuple[str, dict]]:
    if m is None:
        return [("All milestones complete", {"colour": GREY})]
    when = _fmt(m.forecast)
    out = [(f"{m.name}: ", {"colour": INK}), (when, {"bold": True, "colour": INK})]
    if m.forecast and m.baseline and (m.forecast - m.baseline).days > 0:
        out.append((f"  +{(m.forecast - m.baseline).days}d", {"bold": True, "colour": RAG_COLOURS["RED"]}))
    elif m.forecast is None and m.baseline:
        out.append((f"  (was {_fmt(m.baseline)})", {"colour": GREY}))
    return out


# ------------------------------------------------------------------ slide

def build_brief(analysis: PortfolioAnalysis, synthesis: PortfolioSynthesis, ranked: list[RankedPortfolioRisk],
                registry: Registry, out_path: Path, source_pages: int | None = None,
                provider_label: str = "") -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    period = analysis.period_end

    # header -------------------------------------------------------------
    _shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.09, NAVY)
    _text(slide, 0.45, 0.32, 8, 0.5, "AI Transformation Portfolio", size=26, bold=True, colour=NAVY)
    _text(slide, 0.45, 0.82, 8, 0.3,
          f"Leadership brief  ·  status period ending {period.day} {period:%B %Y}  ·  "
          f"{len(analysis.extractions)} initiatives", size=12, colour=GREY)

    counts = analysis.rag_counts()
    _text(slide, 9.2, 0.3, 3.68, 0.25, "ASSESSED STATUS", size=9, bold=True, colour=GREY, align=PP_ALIGN.RIGHT)
    x = 12.88
    for rag in ("GREEN", "AMBER", "RED"):
        w = 1.1
        x -= w
        chip = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.6, w - 0.08, 0.5, RAG_COLOURS[rag])
        chip.adjustments[0] = 0.25
        _label_in(chip, f"{counts[rag]}  {rag.title()}", 13)

    # headline -----------------------------------------------------------
    band = _shape(slide, MSO_SHAPE.RECTANGLE, 0.45, 1.3, 12.43, 0.58, LIGHT)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.45, 1.3, 0.07, 0.58, ACCENT)
    _text(slide, 0.7, 1.3, 12.0, 0.58, synthesis.headline, size=15 if len(synthesis.headline) <= 110 else 13.5,
          bold=True, colour=NAVY,
          anchor=MSO_ANCHOR.MIDDLE)
    del band

    # left panel: initiatives -------------------------------------------
    lx, ly, lw = 0.45, 2.1, 5.95
    _text(slide, lx, ly, lw, 0.3, "Initiative status", size=13, bold=True, colour=NAVY)
    _text(slide, lx + 2.45, ly + 0.05, 1.0, 0.25, "COMPLETE", size=8, bold=True, colour=GREY)
    _text(slide, lx + 3.55, ly + 0.05, 2.4, 0.25, "MILESTONE TO WATCH (FORECAST)", size=8, bold=True, colour=GREY)
    _rule(slide, lx, ly + 0.36, lw)

    ordered = sorted(analysis.extractions.values(), key=lambda e: (e.assessed_rag.rank, e.initiative_id))
    row_h = 3.66 / max(len(ordered), 1)
    for i, e in enumerate(ordered):
        y = ly + 0.45 + i * row_h
        ini = registry[e.initiative_id]
        dot = _shape(slide, MSO_SHAPE.OVAL, lx, y + 0.07, 0.2, 0.2, RAG_COLOURS[e.assessed_rag.value])
        del dot
        _text(slide, lx + 0.3, y + 0.02, 2.1, 0.25, ini.short_name, size=11, bold=True)
        sub = e.initiative_id
        if e.reported_rag != e.assessed_rag:
            sub += f"  ·  reported {e.reported_rag.value.title()}"
        _text(slide, lx + 0.3, y + 0.23, 2.1, 0.2, sub, size=8,
              colour=RAG_COLOURS["RED"] if e.reported_rag != e.assessed_rag else GREY,
              bold=e.reported_rag != e.assessed_rag)
        # % complete bar
        pct = e.percent_complete or 0
        bx = lx + 2.45
        _shape(slide, MSO_SHAPE.RECTANGLE, bx, y + 0.13, 0.62, 0.1, RULE)
        if pct:
            _shape(slide, MSO_SHAPE.RECTANGLE, bx, y + 0.13, 0.62 * pct / 100, 0.1, ACCENT)
        _text(slide, bx + 0.68, y + 0.06, 0.4, 0.25, f"{pct}%" if e.percent_complete is not None else "–",
              size=9, colour=GREY)
        _text(slide, lx + 3.55, y + 0.04, 2.4, row_h - 0.06, milestone_text(watch_milestone(e)), size=9)
        if i < len(ordered) - 1:
            _rule(slide, lx + 0.3, y + row_h - 0.03, lw - 0.3)

    # right panel: risks -------------------------------------------------
    rx, ry, rw = 6.85, 2.1, 6.03
    _text(slide, rx, ry, rw, 0.3, [("Top risks", {"bold": True, "size": 13, "colour": NAVY}),
                                   ("  ranked by impact across the portfolio", {"size": 10, "colour": GREY})])
    _text(slide, rx + rw - 0.7, ry + 0.05, 0.7, 0.25, "SCORE", size=8, bold=True, colour=GREY, align=PP_ALIGN.RIGHT)
    _rule(slide, rx, ry + 0.36, rw)
    top = ranked[:TOP_RISKS]
    rrow = 3.66 / TOP_RISKS
    for i, rp in enumerate(top):
        y = ry + 0.45 + i * rrow
        badge = _shape(slide, MSO_SHAPE.OVAL, rx, y + 0.02, 0.32, 0.32, NAVY if i < 2 else ACCENT)
        _label_in(badge, str(rp.rank), 11)
        _text(slide, rx + 0.45, y, rw - 1.25, 0.25, rp.risk.title, size=11, bold=True)
        _text(slide, rx + 0.45, y + 0.24, rw - 1.25, 0.34, rp.risk.impact_statement, size=8.5, colour=GREY)
        others = [i for i in dict.fromkeys(rp.risk.affected_initiatives) if i != rp.risk.root_initiative]
        hits = f"Root: {rp.risk.root_initiative}" + (f"   ·   Hits: {', '.join(others)}" if others else "")
        _text(slide, rx + 0.45, y + 0.56, rw - 1.25, 0.18, hits, size=8, bold=True, colour=ACCENT)
        score = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, rx + rw - 0.62, y + 0.04, 0.62, 0.32, LIGHT, RULE)
        _label_in(score, f"{rp.portfolio_score:.0f}", 12, colour=NAVY)
        if i < len(top) - 1:
            _rule(slide, rx + 0.45, y + rrow - 0.03, rw - 0.45)

    # decisions strip ----------------------------------------------------
    dy = 6.3
    _text(slide, 0.45, dy - 0.02, 1.6, 0.7, "Decisions\nneeded", size=13, bold=True, colour=NAVY)
    decisions = synthesis.decisions_required[:4]
    dw = (12.43 - 1.55) / max(len(decisions), 1)
    for i, d in enumerate(decisions):
        x = 2.0 + i * dw
        _shape(slide, MSO_SHAPE.RECTANGLE, x, dy, dw - 0.12, 0.72, LIGHT)
        num = _shape(slide, MSO_SHAPE.OVAL, x + 0.1, dy + 0.1, 0.26, 0.26, ACCENT)
        _label_in(num, str(i + 1), 10)
        _text(slide, x + 0.45, dy + 0.08, dw - 0.68, 0.6, d, size=9.5, anchor=MSO_ANCHOR.TOP)

    # footer -------------------------------------------------------------
    src = f"{len(analysis.reports)} status reports"
    if source_pages:
        src += f" ({source_pages} pages)"
    _text(slide, 0.45, 7.12, 12.43, 0.25,
          f"Source: {src}, summarised by LLM{(' (' + provider_label + ')') if provider_label else ''}. "
          "Risk score = likelihood × impact × (1 + 0.25 per other initiative hit). "
          "RAG is the assessed view; where it differs from the team's report, the reported status is shown.",
          size=8, colour=GREY, italic=True)

    slide.notes_slide.notes_text_frame.text = speaker_notes(analysis, synthesis, ranked, registry)

    now = datetime.now()
    props = prs.core_properties
    props.title = f"AI Transformation Portfolio — {period:%d %b %Y}"
    props.subject = "Leadership portfolio brief"
    props.author = props.last_modified_by = "AI Transformation Portfolio Tracker"
    props.created = props.modified = now
    props.revision = 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def speaker_notes(analysis: PortfolioAnalysis, synthesis: PortfolioSynthesis,
                  ranked: list[RankedPortfolioRisk], registry: Registry) -> str:
    lines = ["KEY MESSAGES"]
    lines += [f"{i}. {m}" for i, m in enumerate(synthesis.key_messages, 1)]
    lines += ["", "PORTFOLIO RISKS (ranked)"]
    for rp in ranked:
        lines.append(f"#{rp.rank} [{rp.portfolio_score:.1f}] {rp.risk.title} — root {rp.risk.root_initiative}")
        lines.append(f"    Impact: {rp.risk.impact_statement}")
        lines.append(f"    Action: {rp.risk.recommended_action}")
        lines.append(f"    Rolls up: {', '.join(rp.risk.linked_risk_ids)}")
    lines += ["", "RAG CHANGES (reported -> assessed)"]
    changed = [e for e in analysis.extractions.values() if e.reported_rag != e.assessed_rag]
    lines += [f"{e.initiative_id}: {e.reported_rag.value} -> {e.assessed_rag.value}. {e.rag_rationale}"
              for e in changed] or ["None"]
    if analysis.warnings:
        lines += ["", "DATA QUALITY WARNINGS"] + [f"- {w}" for w in analysis.warnings]
    return "\n".join(lines)
