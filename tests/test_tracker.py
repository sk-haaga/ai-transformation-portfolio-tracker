from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from pptx import Presentation

from portfolio_tracker.analysis import analyse, is_grounded, rank_portfolio_risks
from portfolio_tracker.brief import watch_milestone
from portfolio_tracker.ingest import load_reports
from portfolio_tracker.llm import ClaudeProvider, ReplayProvider, inline_refs
from portfolio_tracker.pipeline import run
from portfolio_tracker.registry import Registry
from portfolio_tracker.schema import InitiativeExtraction, Milestone, PortfolioSynthesis

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "data" / "status_reports"
REFERENCE = ROOT / "data" / "reference_run"


@pytest.fixture(scope="module")
def registry() -> Registry:
    return Registry.from_csv(ROOT / "data" / "initiatives.csv")


@pytest.fixture(scope="module")
def reference(registry):
    reports = load_reports(REPORTS, registry)
    provider = ReplayProvider(REFERENCE)
    extractions = [provider.extract(r, registry) for r in reports]
    return reports, extractions, provider


# ---------------------------------------------------------------- registry

@pytest.mark.parametrize("mention, expected", [
    ("DAT-02", "DAT-02"),
    ("the Identity programme (IT-07)", "IT-07"),
    ("Data Platform team", "DAT-02"),
    ("AI Skills Academy", "HR-05"),
    ("Customer Service Copilot", "CRM-01"),
    ("Procurement", None),
    ("Legal", None),
])
def test_registry_resolves_loose_mentions(registry, mention, expected):
    assert registry.resolve(mention) == expected


def test_reports_are_mapped_to_initiatives(registry):
    reports = load_reports(REPORTS, registry)
    assert sorted(r.initiative_id for r in reports) == sorted(registry.ids())


# ---------------------------------------------------------------- grounding

def test_grounding_accepts_verbatim_quote_across_markdown_and_whitespace():
    src = "**Governance approval delay (HIGH).** No checklist = no approval =\n no go-live. Escalated."
    assert is_grounded("No checklist = no approval = no go-live.", src)


def test_grounding_rejects_invented_quote():
    src = "The platform is stable and the finance domain is live."
    assert not is_grounded("The platform has been decommissioned ahead of plan", src)


def test_every_reference_risk_is_grounded(reference, registry):
    reports, extractions, _ = reference
    analysis = analyse(reports, extractions, registry)
    assert all(r.grounded for r in analysis.scored_risks)
    assert analysis.warnings == []


def test_fabricated_evidence_is_flagged(reference, registry):
    reports, extractions, _ = reference
    bad = extractions[0].model_copy(deep=True)
    bad.risks[0].evidence = "the vendor has gone bankrupt and support has ended"
    analysis = analyse(reports, [bad, *extractions[1:]], registry)
    flagged = [r for r in analysis.scored_risks if not r.grounded]
    assert [r.risk_id for r in flagged] == [f"{bad.initiative_id}.R1"]
    assert any("evidence not found" in w for w in analysis.warnings)


# ---------------------------------------------------------------- analysis

def test_dependency_reach(reference, registry):
    reports, extractions, _ = reference
    a = analyse(reports, extractions, registry)
    # three initiatives are blocked or at risk because of the data platform and governance
    assert a.direct_reach["DAT-02"] == 3
    assert a.direct_reach["GOV-03"] == 3
    assert a.direct_reach["FIN-04"] == 0
    # FIN-04's dependency on IT-07 is on track, so it does not count as reach
    assert a.direct_reach["IT-07"] == 1


def test_risk_scores_are_likelihood_times_impact_weighted_by_reach(reference, registry):
    reports, extractions, _ = reference
    a = analyse(reports, extractions, registry)
    r = a.risk_index["DAT-02.R3"]
    assert r.score == r.risk.likelihood * r.risk.impact == 25
    assert r.portfolio_score == pytest.approx(25 * (1 + 0.25 * 3))


def test_portfolio_ranking_is_computed_not_taken_from_llm(reference, registry):
    reports, extractions, provider = reference
    a = analyse(reports, extractions, registry)
    synthesis = provider.synthesise("", "", "", "")
    reversed_synthesis = synthesis.model_copy(update={"portfolio_risks": synthesis.portfolio_risks[::-1]})
    ranked = rank_portfolio_risks(synthesis, a, registry)
    ranked_rev = rank_portfolio_risks(reversed_synthesis, a, registry)
    assert [r.risk.title for r in ranked] == [r.risk.title for r in ranked_rev]
    assert [r.portfolio_score for r in ranked] == sorted((r.portfolio_score for r in ranked), reverse=True)
    assert ranked[0].risk.root_initiative == "DAT-02"


def test_unknown_risk_ids_in_synthesis_are_warned(reference, registry):
    reports, extractions, provider = reference
    a = analyse(reports, extractions, registry)
    s = provider.synthesise("", "", "", "")
    s.portfolio_risks[0].linked_risk_ids.append("XYZ-99.R7")
    rank_portfolio_risks(s, a, registry)
    assert any("XYZ-99.R7" in w for w in a.warnings)


def test_watch_milestone_prefers_largest_slip():
    e = SimpleNamespace(milestones=[
        Milestone(name="small", baseline=date(2026, 10, 1), forecast=date(2026, 10, 8), status="late"),
        Milestone(name="big", baseline=date(2026, 11, 1), forecast=date(2027, 1, 1), status="at_risk"),
        Milestone(name="done", baseline=date(2026, 9, 1), forecast=date(2026, 12, 1), status="done"),
    ])
    assert watch_milestone(e).name == "big"


# ---------------------------------------------------------------- LLM layer

def test_tool_schemas_have_no_refs():
    for model in (InitiativeExtraction, PortfolioSynthesis):
        assert "$ref" not in json.dumps(inline_refs(model.model_json_schema()))


def test_claude_provider_retries_once_on_invalid_output(reference, registry, tmp_path):
    reports, extractions, _ = reference
    good = json.loads(extractions[0].model_dump_json())
    bad = {**good, "percent_complete": 140}

    calls = []

    def create(**kwargs):
        calls.append(kwargs)
        payload = bad if len(calls) == 1 else good
        return SimpleNamespace(content=[SimpleNamespace(type="tool_use", id=f"t{len(calls)}", input=payload)])

    provider = object.__new__(ClaudeProvider)
    provider.client = SimpleNamespace(messages=SimpleNamespace(create=create))
    provider.model, provider.max_tokens, provider.cache_dir = "test-model", 1000, tmp_path

    result = provider.extract(reports[0], registry)
    assert result.percent_complete == good["percent_complete"]
    assert len(calls) == 2
    assert calls[0]["tool_choice"] == {"type": "tool", "name": "record_initiative_status"}
    assert calls[1]["messages"][-1]["content"][0]["is_error"] is True

    # second run is served from the cache, no new API call
    provider.extract(reports[0], registry)
    assert len(calls) == 2


# ---------------------------------------------------------------- end to end

def test_pipeline_end_to_end(tmp_path):
    res = run(REPORTS, ROOT / "data" / "initiatives.csv", ReplayProvider(REFERENCE), tmp_path,
              save_run_dir=tmp_path / "replay")

    prs = Presentation(res.brief_path)
    assert len(prs.slides) == 1
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "KEY MESSAGES" in notes and "PORTFOLIO RISKS" in notes
    slide_text = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert res.ranked[0].risk.title in slide_text

    assert res.source_pages > 5
    for table in ("dim_initiative", "fact_milestone", "fact_risk", "fact_dependency",
                  "portfolio_risk", "bridge_portfolio_risk", "run_metadata"):
        assert (res.powerbi_dir / f"{table}.csv").exists()
    dim = pd.read_csv(res.powerbi_dir / "dim_initiative.csv")
    assert len(dim) == 8 and set(dim["assessed_rag"]) <= {"RED", "AMBER", "GREEN"}
    assert res.roadmap_path.exists()

    # a saved run replays to the same result
    again = run(REPORTS, ROOT / "data" / "initiatives.csv", ReplayProvider(tmp_path / "replay"), tmp_path / "2")
    assert [r.risk.title for r in again.ranked] == [r.risk.title for r in res.ranked]
